from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

WORK = Path(__file__).parent
REFERENCE = WORK / "reference.docx"
OUTPUT = Path(__file__).parents[1] / "Bao cao BTL Smart Garden ESP32 DH13C6.docx"
LOGO = WORK / "hunre-logo.jpeg"


def set_font(run, name="Times New Roman", size=13, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6", val="single"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{key}"))
        if element is None:
            element = OxmlElement(f"w:{key}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_field(paragraph, instruction, placeholder=""):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = placeholder
    text_run.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    begin_run = OxmlElement("w:r")
    begin_run.append(begin)
    instr_run = OxmlElement("w:r")
    instr_run.append(instr)
    separate_run = OxmlElement("w:r")
    separate_run.append(separate)
    end_run = OxmlElement("w:r")
    end_run.append(end)
    paragraph._p.extend([begin_run, instr_run, separate_run, text_run, end_run])


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.first_line_indent = Cm(1.27)
    normal.paragraph_format.space_after = Pt(4)

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True

    if "Chapter Title" not in [style.name for style in doc.styles]:
        chapter_title = doc.styles.add_style("Chapter Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        chapter_title = doc.styles["Chapter Title"]
    chapter_title.base_style = title
    chapter_title.font.name = "Times New Roman"
    chapter_title._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    chapter_title._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    chapter_title._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    chapter_title.font.size = Pt(16)
    chapter_title.font.bold = True
    chapter_title.font.color.rgb = RGBColor(0, 0, 0)
    chapter_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chapter_title.paragraph_format.space_after = Pt(8)
    chapter_title.paragraph_format.keep_with_next = True
    chapter_ppr = chapter_title._element.get_or_add_pPr()
    outline = chapter_ppr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        chapter_ppr.append(outline)
    outline.set(qn("w:val"), "0")

    heading_specs = {
        "Heading 1": (14, True, False, 18, 5),
        "Heading 2": (13, True, False, 10, 4),
        "Heading 3": (13, True, True, 8, 3),
        "Heading 4": (13, False, True, 5, 2),
    }
    for name, (size, bold, italic, before, after) in heading_specs.items():
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style_ppr = style._element.get_or_add_pPr()
        inherited_numbering = style_ppr.find(qn("w:numPr"))
        if inherited_numbering is not None:
            style_ppr.remove(inherited_numbering)

    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    caption.font.size = Pt(12)
    caption.font.bold = True
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False
    caption.paragraph_format.keep_together = True


def clear_document_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_page_geometry(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)


def add_page_number(section):
    footer = section.footer
    for existing in list(footer.paragraphs):
        existing._element.getparent().remove(existing._element)
    paragraph = footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    add_field(paragraph, " PAGE ", "1")
    for run in paragraph.runs:
        set_font(run, size=11)


def add_page_border(section, color="000000", size="10", space="18"):
    sect_pr = section._sectPr
    pg_borders = sect_pr.find(qn("w:pgBorders"))
    if pg_borders is None:
        pg_borders = OxmlElement("w:pgBorders")
        sect_pr.append(pg_borders)
    pg_borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), space)
        element.set(qn("w:color"), color)
        pg_borders.append(element)


def add_cover_line(doc, text, size=13, bold=True, space_after=2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(space_after)
    set_font(p.add_run(text), size=size, bold=bold)
    return p


def add_body(doc, text, bold_lead=None, italic=False, no_indent=False):
    p = doc.add_paragraph(style="Normal")
    if no_indent:
        p.paragraph_format.first_line_indent = None
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), bold=True)
        set_font(p.add_run(text[len(bold_lead):]), italic=italic)
    else:
        set_font(p.add_run(text), italic=italic)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Paragraph")
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run("- "))
        set_font(p.add_run(item))


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.first_line_indent = None
    for run in p.runs:
        set_font(run, size=14 if level == 1 else 13, bold=True, italic=level >= 3)
    return p


def add_chapter(doc, title):
    doc.add_page_break()
    p = doc.add_paragraph(title, style="Chapter Title")
    p.paragraph_format.first_line_indent = None
    for run in p.runs:
        set_font(run, size=16, bold=True)
    return p


def add_caption(doc, label, text):
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.first_line_indent = None
    set_font(p.add_run(f"{label} "), size=12, bold=True, italic=True)
    add_field(p, f" SEQ {label} \\* ARABIC ", "1")
    set_font(p.add_run(f" {text}"), size=12, bold=True, italic=True)
    return p


def add_placeholder(doc, label, caption, lines=7):
    outer = doc.add_table(rows=1, cols=1)
    outer.alignment = WD_TABLE_ALIGNMENT.CENTER
    outer.autofit = False
    outer_cell = outer.cell(0, 0)
    outer_cell.width = Cm(14.8)
    set_cell_border(outer_cell, color="FFFFFF", size="0", val="nil")
    set_cell_margins(outer_cell, top=0, start=0, bottom=0, end=0)
    outer_cell.paragraphs[0].paragraph_format.space_after = Pt(0)

    image_table = outer_cell.add_table(rows=1, cols=1)
    image_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    image_table.autofit = False
    cell = image_table.cell(0, 0)
    cell.width = Cm(14.8)
    set_cell_border(cell, color="A6A6A6", size="8", val="dashed")
    set_cell_margins(cell, top=180, start=180, bottom=180, end=180)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(f"[CHÈN ẢNH: {caption.upper()}]"), size=11, bold=True, color="7F7F7F")
    for _ in range(max(2, lines - 2)):
        spacer = cell.add_paragraph()
        spacer.paragraph_format.first_line_indent = None
        spacer.paragraph_format.space_after = Pt(5)
        spacer.add_run(" ")
    note = cell.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.first_line_indent = None
    set_font(note.add_run("Thay khung này bằng ảnh thực tế, giữ nguyên chú thích bên dưới"), size=10, italic=True, color="7F7F7F")

    caption_p = outer_cell.add_paragraph()
    caption_p.style = doc.styles["Caption"]
    caption_p.paragraph_format.first_line_indent = None
    set_font(caption_p.add_run(f"{label} "), size=12, bold=True, italic=True)
    add_field(caption_p, f" SEQ {label} \\* ARABIC ", "1")
    set_font(caption_p.add_run(f" {caption}"), size=12, bold=True, italic=True)
    outer_tr_pr = outer.rows[0]._tr.get_or_add_trPr()
    outer_tr_pr.append(OxmlElement("w:cantSplit"))
    inner_tr_pr = image_table.rows[0]._tr.get_or_add_trPr()
    inner_tr_pr.append(OxmlElement("w:cantSplit"))


def add_table(doc, label, caption, headers, rows, widths=None, small=False, keep_together=False):
    cap = add_caption(doc, label, caption)
    cap.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_row = table.rows[0]
    repeat_table_header(header_row)
    for i, header in enumerate(headers):
        cell = header_row.cells[i]
        if widths:
            cell.width = Cm(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        if small:
            set_cell_margins(cell, top=80, start=60, bottom=80, end=60)
        else:
            set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(str(header)), size=9.5 if small else 11, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            if widths:
                cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if small:
                set_cell_margins(cell, top=80, start=60, bottom=80, end=60)
            else:
                set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, "EAF2F8")
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or len(value) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(str(value)), size=9.5 if small else 11)
    if keep_together:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.line_spacing = Pt(1)
    after.paragraph_format.first_line_indent = Cm(0)
    set_font(after.add_run(" "), size=1)
    return table


def add_code(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_font(p.add_run(line), name="Courier New", size=9.5)


def add_toc_field(doc, instruction, placeholder):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = 1.15
    add_field(p, instruction, placeholder)


def extract_logo():
    if LOGO.exists():
        return
    with ZipFile(REFERENCE) as archive:
        data = archive.read("word/media/image1.jpeg")
    LOGO.write_bytes(data)


def main():
    extract_logo()
    doc = Document(REFERENCE)
    clear_document_body(doc)
    configure_styles(doc)
    first = doc.sections[0]
    set_page_geometry(first)
    add_page_border(first)
    for existing in list(first.footer.paragraphs):
        existing._element.getparent().remove(existing._element)
    first.footer.add_paragraph()

    add_cover_line(doc, "TRƯỜNG ĐẠI HỌC TÀI NGUYÊN VÀ MÔI TRƯỜNG HÀ NỘI", 13)
    add_cover_line(doc, "KHOA CÔNG NGHỆ THÔNG TIN", 13)
    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.paragraph_format.first_line_indent = None
    logo_p.paragraph_format.space_before = Pt(10)
    logo_p.paragraph_format.space_after = Pt(10)
    logo_p.add_run().add_picture(str(LOGO), width=Cm(4.8))
    add_cover_line(doc, "BÁO CÁO BÀI TẬP LỚN", 16, True, 10)
    title_p = doc.add_paragraph(style="Title")
    title_p.paragraph_format.first_line_indent = None
    title_p.paragraph_format.space_after = Pt(8)
    set_font(title_p.add_run("XÂY DỰNG MÔ HÌNH VƯỜN THÔNG MINH CÓ CẢM BIẾN ĐO ĐỘ ẨM ĐẤT, NHIỆT ĐỘ ĐỘ ẨM CÓ SỬ DỤNG ESP32"), size=16, bold=True)
    add_cover_line(doc, "HỌC PHẦN: PHÁT TRIỂN ỨNG DỤNG HỆ THỐNG NHÚNG VÀ IoT", 13, True, 14)

    cover_info = doc.add_table(rows=4, cols=2)
    cover_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_info.autofit = False
    info = [
        ("Giảng viên hướng dẫn", "Nguyễn Thành Long"),
        ("Lớp", "DH13C6"),
        ("Nhóm thực hiện", "Nhóm 6"),
        ("Sinh viên thực hiện", ""),
    ]
    for r, (left, right) in enumerate(info):
        cover_info.cell(r, 0).width = Cm(6.0)
        cover_info.cell(r, 1).width = Cm(7.0)
        for c, value in enumerate((left, right)):
            cell = cover_info.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell._tc.get_or_add_tcPr().append(OxmlElement("w:tcBorders"))
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run((value + " :") if c == 0 else value), size=12, bold=c == 0)

    doc.add_paragraph()
    member_rows = [["", "", "DH13C6"] for _ in range(5)]
    add_table(doc, "Bảng", "Danh sách thành viên thực hiện đề tài", ["Mã sinh viên", "Họ và tên", "Lớp"], member_rows, [4.0, 7.2, 3.0], small=True)
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.paragraph_format.first_line_indent = None
    date_p.paragraph_format.space_before = Pt(10)
    set_font(date_p.add_run("HÀ NỘI - 2026"), size=13, bold=True)

    second = doc.add_section(WD_SECTION.NEW_PAGE)
    set_page_geometry(second)
    second.footer.is_linked_to_previous = False
    page_numbering = second._sectPr.find(qn("w:pgNumType"))
    if page_numbering is None:
        page_numbering = OxmlElement("w:pgNumType")
        second._sectPr.append(page_numbering)
    page_numbering.set(qn("w:start"), "1")
    add_page_number(second)

    doc.add_paragraph("MỤC LỤC", style="Title")
    add_toc_field(doc, ' TOC \\t "Chapter Title,1,Heading 1,2,Heading 2,3" \\h \\z ', "Nhấn Ctrl+A, F9 trong Word nếu mục lục chưa cập nhật")
    doc.add_page_break()
    doc.add_paragraph("DANH MỤC HÌNH ẢNH", style="Title")
    add_toc_field(doc, ' TOC \\h \\z \\c "Hình" ', "Danh mục hình sẽ được cập nhật khi mở trong Word")
    doc.add_page_break()
    doc.add_paragraph("DANH MỤC BẢNG", style="Title")
    add_toc_field(doc, ' TOC \\h \\z \\c "Bảng" ', "Danh mục bảng sẽ được cập nhật khi mở trong Word")

    add_chapter(doc, "MỞ ĐẦU")
    add_heading(doc, "1 Đặt vấn đề", 1)
    add_body(doc, "Tưới nước đúng thời điểm là một yêu cầu quan trọng trong chăm sóc cây trồng. Trong mô hình vườn nhỏ, chậu cây hoặc khu trồng thử nghiệm, người dùng thường đánh giá độ ẩm bằng cảm giác hoặc tưới theo lịch cố định. Cách làm này khó phản ánh chính xác trạng thái của đất, dễ dẫn đến tưới thiếu khi thời tiết khô nóng hoặc tưới thừa khi đất vẫn còn ẩm.")
    add_body(doc, "Đề tài xây dựng một mô hình vườn thông minh dùng ESP32 làm bộ điều khiển trung tâm. Hệ thống thu nhận độ ẩm đất, nhiệt độ và độ ẩm không khí; hiển thị dữ liệu trên giao diện web; cho phép điều khiển máy bơm nước MB370 bằng tay hoặc theo ngưỡng tự động. Dữ liệu thời tiết OpenWeather và chatbot Gemini được tích hợp để hỗ trợ người dùng đánh giá điều kiện môi trường và tham khảo cấu hình tưới.")
    add_body(doc, "Sản phẩm tập trung vào một mô hình học tập có đầy đủ chuỗi xử lý của hệ thống nhúng và IoT: cảm biến tạo dữ liệu, ESP32 xử lý và cung cấp API trong mạng Wi-Fi, ứng dụng web trực quan hóa trạng thái, relay đóng cắt tải, còn máy bơm thực hiện tưới. Thiết kế ưu tiên khả năng quan sát, dễ kiểm tra và có cơ chế an toàn để hạn chế bơm hoạt động kéo dài.")

    add_heading(doc, "2 Mục tiêu đề tài", 1)
    add_bullets(doc, [
        "Đo độ ẩm đất bằng cảm biến Soil Moisture và quy đổi giá trị ADC thành phần trăm sau khi hiệu chuẩn.",
        "Đo nhiệt độ và độ ẩm không khí bằng DHT11, đồng thời phát hiện trạng thái đọc lỗi.",
        "Điều khiển máy bơm nước MB370 thông qua relay 2 kênh ở chế độ thủ công và tự động.",
        "Truyền dữ liệu từ ESP32 tới ứng dụng web qua mạng Wi-Fi nội bộ bằng API JSON.",
        "Hiển thị cảm biến, trạng thái bơm, cảnh báo, thời tiết khu vực và dự báo trên giao diện web đáp ứng máy tính và điện thoại.",
        "Tích hợp chatbot sử dụng dữ liệu cảm biến và tên cây để đề xuất ngưỡng tưới; người dùng được quyền áp dụng, xem xét hoặc từ chối.",
        "Xây dựng cơ chế kiểm tra đầu vào, timeout, giới hạn thời gian tưới và phản hồi lỗi rõ ràng.",
    ])

    add_heading(doc, "3 Phạm vi và sản phẩm dự kiến", 1)
    add_body(doc, "Phạm vi phần cứng gồm ESP32, DHT11, cảm biến độ ẩm đất, máy bơm nước MB370, mạch mở rộng, bộ chuyển đổi nguồn, relay 2 kênh và các phụ kiện kết nối. Phạm vi phần mềm gồm firmware C/C++ cho ESP32 và ứng dụng web React chạy trong cùng mạng cục bộ. Hệ thống chưa triển khai điều khiển nhiều vùng tưới, chưa có cảm biến lưu lượng nước và chưa đồng bộ dữ liệu dài hạn lên máy chủ đám mây.")
    add_body(doc, "Sản phẩm bàn giao gồm mô hình phần cứng, firmware ESP32, ứng dụng web Smart Garden và báo cáo kỹ thuật. Các ảnh thực tế sẽ được bổ sung vào đúng các vị trí đã đánh dấu trong tài liệu này sau khi nhóm hoàn thiện lắp ráp và chụp thử nghiệm.")

    add_table(doc, "Bảng", "Các sản phẩm của đề tài", ["STT", "Sản phẩm", "Nội dung"], [
        ["1", "Mô hình phần cứng", "ESP32, cảm biến, relay và máy bơm được lắp thành một hệ thống tưới thử nghiệm."],
        ["2", "Firmware ESP32", "Đọc cảm biến, điều khiển relay, xử lý tưới tự động và cung cấp API JSON."],
        ["3", "Ứng dụng web", "Theo dõi thời tiết, cảm biến, điều khiển bơm, cấu hình tự động và chatbot."],
        ["4", "Báo cáo", "Mô tả yêu cầu, thiết kế, triển khai, kiểm thử, kết quả và hướng phát triển."],
    ], [1.3, 4.0, 9.0])

    add_chapter(doc, "CHƯƠNG 1 TỔNG QUAN VẤN ĐỀ NGHIÊN CỨU")
    add_heading(doc, "1 Lý do chọn đề tài", 1)
    add_body(doc, "Nhu cầu chăm cây trong gia đình và mô hình trồng thử nghiệm ngày càng phổ biến, nhưng việc tưới nước vẫn phụ thuộc nhiều vào kinh nghiệm. Đất trong các chậu có thể khô nhanh vào ngày nắng, trong khi độ ẩm cao hoặc mưa làm nhu cầu tưới giảm. Một hệ thống tự động dựa trên số đo giúp người dùng phản ứng theo trạng thái thực tế thay vì chỉ dựa vào lịch cố định.")
    add_body(doc, "ESP32 phù hợp với bài toán vì tích hợp Wi-Fi, có các chân ADC để đọc cảm biến đất và đủ GPIO để kết nối DHT11 cùng relay. Việc sử dụng giao diện web giúp thiết bị không phụ thuộc vào một ứng dụng cài đặt riêng; máy tính hoặc điện thoại trong cùng mạng có thể truy cập, quan sát và gửi lệnh. Đây cũng là đề tài kết hợp rõ ràng giữa phần cứng, firmware, giao tiếp mạng và phát triển giao diện người dùng.")

    add_heading(doc, "2 Ý nghĩa thực tiễn", 1)
    add_body(doc, "Mô hình giúp giảm thao tác kiểm tra đất và bật bơm lặp lại. Người dùng có thể xem độ ẩm đất, nhiệt độ và độ ẩm không khí ngay trên trình duyệt, từ đó quyết định tưới hoặc kích hoạt chế độ tự động. Khi được hiệu chuẩn đúng và đặt ngưỡng phù hợp, hệ thống hỗ trợ duy trì độ ẩm ổn định hơn cho cây.")
    add_body(doc, "Về mặt học tập, đề tài cho phép nhóm thực hành các nội dung cốt lõi của hệ thống nhúng và IoT: đọc tín hiệu analog, xử lý cảm biến số, điều khiển tải qua relay, thiết kế API, quản lý trạng thái thời gian thực, xử lý lỗi mạng, xây dựng giao diện và tích hợp dịch vụ dữ liệu bên ngoài. Những kiến thức này có thể mở rộng sang nhà kính, tưới nhiều khu vực hoặc giám sát từ xa.")

    add_heading(doc, "3 Tổng quan hệ thống nhúng và IoT", 1)
    add_body(doc, "Hệ thống nhúng là tổ hợp phần cứng và phần mềm được thiết kế để thực hiện một nhóm chức năng cụ thể. Trong đề tài, ESP32 thực hiện việc đọc cảm biến, quyết định bật hoặc tắt bơm và phản hồi yêu cầu từ ứng dụng web. Chu kỳ xử lý được thiết kế không chặn trong phần đọc định kỳ để máy chủ web vẫn có thể phục vụ yêu cầu điều khiển.")
    add_body(doc, "Khía cạnh IoT thể hiện ở khả năng kết nối ESP32 với mạng Wi-Fi và trao đổi dữ liệu với trình duyệt. Thiết bị cung cấp các endpoint HTTP trả về JSON; ứng dụng web thăm dò dữ liệu định kỳ, hiển thị trạng thái và gửi cấu hình. Trong phiên bản hiện tại, kết nối hoạt động trong mạng LAN để thuận tiện cho mô hình học tập.")

    add_heading(doc, "4 Thành phần phần cứng", 1)
    add_heading(doc, "4.1 ESP32", 2)
    add_body(doc, "ESP32 là bộ điều khiển trung tâm, đảm nhiệm đọc dữ liệu, chạy thuật toán tưới và giao tiếp Wi-Fi. Firmware của dự án sử dụng GPIO34 thuộc nhóm ADC1 để đọc cảm biến độ ẩm đất; lựa chọn này phù hợp khi Wi-Fi hoạt động vì ADC2 có các hạn chế dùng chung tài nguyên với Wi-Fi. GPIO4 nhận dữ liệu DHT11 và GPIO26 điều khiển relay.")
    add_table(doc, "Bảng", "Vai trò và cấu hình ESP32 trong dự án", ["Nội dung", "Giá trị"], [
        ["Kết nối mạng", "Wi-Fi chế độ Station, mDNS smartgarden.local"],
        ["Chân DHT11", "GPIO4"],
        ["Chân cảm biến đất", "GPIO34 - ADC1, chỉ dùng làm đầu vào"],
        ["Chân relay", "GPIO26"],
        ["Chu kỳ đọc cảm biến", "2 giây trên firmware"],
        ["Máy chủ HTTP", "Cổng 80, dữ liệu JSON"],
    ], [5.1, 9.2])
    add_placeholder(doc, "Hình", "Board ESP32 sử dụng trong mô hình", 6)

    add_heading(doc, "4.2 Cảm biến nhiệt độ và độ ẩm DHT11", 2)
    add_body(doc, "DHT11 cung cấp dữ liệu nhiệt độ và độ ẩm không khí qua tín hiệu số một dây. Cảm biến phù hợp cho mô hình học tập và giám sát xu hướng môi trường, nhưng độ chính xác và tốc độ phản hồi thấp hơn các cảm biến cao cấp. Firmware kiểm tra giá trị NaN; khi đọc lỗi, API trả về dht11_ok bằng false để giao diện không hiển thị số liệu sai.")
    add_placeholder(doc, "Hình", "Cảm biến DHT11 và vị trí lắp trên mô hình", 5)

    add_heading(doc, "4.3 Cảm biến độ ẩm đất Soil Moisture", 2)
    add_body(doc, "Cảm biến Soil Moisture tạo tín hiệu analog thay đổi theo độ dẫn điện của đất. Firmware đọc giá trị ADC tại GPIO34 và ánh xạ từ hai mốc hiệu chuẩn khô và ướt sang phần trăm. Trong mã hiện tại, mốc khô là 3000 và mốc ướt là 1200; hai giá trị này phải được đo lại với chính cảm biến, loại đất và nguồn cấp của mô hình trước khi dùng cho đánh giá cuối cùng.")
    add_body(doc, "Cảm biến điện trở có thể bị ăn mòn khi cắm lâu trong đất ẩm. Khi vận hành dài hạn, cần hạn chế cấp nguồn liên tục cho đầu dò hoặc thay bằng cảm biến điện dung. Đầu dò cần đặt ở vùng rễ và tránh ngay sát vòi tưới để giá trị đại diện tốt hơn cho toàn bộ chậu.")
    add_placeholder(doc, "Hình", "Cảm biến độ ẩm đất Soil Moisture và đầu dò", 5)

    add_heading(doc, "4.4 Máy bơm nước MB370", 2)
    add_body(doc, "Máy bơm nước MB370 là cơ cấu chấp hành đưa nước từ bình chứa tới cây. Bơm là tải DC có dòng khởi động lớn hơn khả năng cấp trực tiếp của chân GPIO, vì vậy phải dùng nguồn phù hợp và đóng cắt qua relay. Điện áp và dòng định mức cần đối chiếu trên nhãn của bơm thực tế trước khi chọn bộ chuyển đổi nguồn.")
    add_body(doc, "Đường ống cần được cố định, đầu hút phải ngập nước và đầu ra không được gập. Khi kiểm thử, thời gian tưới nên bắt đầu từ mức ngắn, sau đó đo lượng nước thực tế theo giây để hiệu chỉnh cấu hình phù hợp với loại cây và kích thước chậu.")
    add_placeholder(doc, "Hình", "Máy bơm nước MB370 và đường ống tưới", 5)

    add_heading(doc, "4.5 Relay 2 kênh", 2)
    add_body(doc, "Relay 2 kênh cách ly tín hiệu điều khiển logic khỏi tải bơm. Firmware hiện sử dụng một kênh tại GPIO26 và điều khiển theo mức active LOW: xuất LOW để bật, HIGH để tắt. Kênh còn lại có thể dự phòng cho van điện từ hoặc bơm thứ hai trong phiên bản mở rộng.")
    add_body(doc, "Tiếp điểm COM, NO và NC phải đấu đúng theo yêu cầu an toàn. Với cấu hình thông thường, bơm nối qua COM và NO để mặc định tắt khi relay không được kích. Nguồn bơm không lấy trực tiếp từ chân 3.3 V của ESP32; dây nguồn và tiếp điểm phải chịu được dòng của bơm.")
    add_placeholder(doc, "Hình", "Relay 2 kênh và cách đấu máy bơm", 5)

    add_heading(doc, "4.6 Mạch mở rộng và bộ chuyển đổi", 2)
    add_body(doc, "Mạch mở rộng được dùng để đưa các chân ESP32 ra hàng cắm thuận tiện, phân phối nguồn và giảm lỗi nối dây khi lắp ráp. Bộ chuyển đổi nguồn tạo mức điện áp phù hợp cho ESP32, relay và máy bơm từ nguồn đầu vào của mô hình. Trước khi cấp điện, nhóm cần đo điện áp đầu ra bằng đồng hồ và kiểm tra cực tính.")
    add_body(doc, "Thiết kế nguồn nên tách nhánh tải bơm khỏi nhánh logic, đồng thời nối chung điểm tham chiếu khi module điều khiển yêu cầu. Có thể bổ sung cầu chì, công tắc nguồn và tụ lọc gần relay hoặc bơm để giảm sụt áp và nhiễu. Mọi thông số đầu vào, đầu ra và dòng tối đa phải được ghi theo linh kiện thực tế sau khi lắp.")
    add_placeholder(doc, "Hình", "Mạch mở rộng và bộ chuyển đổi nguồn của hệ thống", 5)

    add_table(doc, "Bảng", "Danh sách phần cứng chính", ["STT", "Thiết bị", "Số lượng", "Chức năng"], [
        ["1", "ESP32", "1", "Xử lý trung tâm, Wi-Fi, API và thuật toán điều khiển."],
        ["2", "DHT11", "1", "Đo nhiệt độ và độ ẩm không khí."],
        ["3", "Soil Moisture", "1", "Đo độ ẩm đất qua tín hiệu analog."],
        ["4", "Máy bơm nước MB370", "1", "Bơm nước tưới cây."],
        ["5", "Relay 2 kênh", "1", "Đóng cắt nguồn cho bơm, dự phòng một kênh."],
        ["6", "Mạch mở rộng", "1", "Phân phối chân kết nối và nguồn."],
        ["7", "Bộ chuyển đổi nguồn", "1", "Tạo điện áp phù hợp cho các nhánh tải."],
        ["8", "Nguồn, dây, ống nước", "Theo mô hình", "Cấp nguồn và dẫn nước."],
    ], [1.2, 4.0, 2.0, 7.1], small=True, keep_together=True)

    add_chapter(doc, "CHƯƠNG 2 PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG")
    add_heading(doc, "1 Phân tích yêu cầu", 1)
    add_heading(doc, "1.1 Yêu cầu chức năng", 2)
    add_bullets(doc, [
        "Kết nối tới ESP32 bằng địa chỉ IP hoặc tên mDNS trong cùng mạng Wi-Fi.",
        "Đọc và hiển thị độ ẩm đất, nhiệt độ, độ ẩm không khí, trạng thái relay và cấu hình tưới.",
        "Cho phép bật hoặc tắt bơm bằng tay với thời gian giới hạn khi chế độ tự động đang tắt.",
        "Cho phép bật chế độ tưới tự động, đặt ngưỡng bắt đầu, ngưỡng dừng và thời gian tối đa.",
        "Tìm kiếm thời tiết khu vực và hiển thị dự báo, cảnh báo liên quan bên phần ESP32.",
        "Cho phép nhập tên cây và chuyển số đo hiện tại tới chatbot để tham khảo cấu hình tưới.",
        "Không tự động áp dụng gợi ý AI; người dùng phải chọn áp dụng, xem xét hoặc từ chối.",
    ])
    add_heading(doc, "1.2 Yêu cầu phi chức năng", 2)
    add_bullets(doc, [
        "Giao diện rõ ràng, chỉ gồm hai khu vực chính và thích nghi với màn hình nhỏ.",
        "Phản hồi lỗi kết nối, timeout, dữ liệu thiếu và cấu hình không hợp lệ bằng tiếng Việt dễ hiểu.",
        "Không để bơm chạy quá 10 phút liên tục trong mọi chế độ.",
        "Chặn cấu hình có ngưỡng bật lớn hơn hoặc bằng ngưỡng dừng và thời gian ngoài 1 đến 600 giây.",
        "Không sử dụng dữ liệu DHT11 khi cảm biến báo lỗi; không áp dụng gợi ý cũ sau khi đổi thiết bị hoặc tên cây.",
        "Các khóa API phải được lưu trong tệp môi trường và không ghi trực tiếp vào mã nguồn công khai.",
    ])

    add_heading(doc, "2 Kiến trúc tổng thể", 1)
    add_body(doc, "Hệ thống gồm bốn lớp. Lớp cảm nhận gồm DHT11 và cảm biến đất. Lớp điều khiển là ESP32 cùng firmware. Lớp chấp hành gồm relay và máy bơm. Lớp ứng dụng là giao diện web nhận dữ liệu từ ESP32, OpenWeather và Gemini. Dữ liệu từ cảm biến đi lên giao diện; lệnh điều khiển và cấu hình đi từ giao diện xuống ESP32.")
    add_placeholder(doc, "Hình", "Sơ đồ khối tổng thể của mô hình vườn thông minh", 8)
    add_table(doc, "Bảng", "Luồng dữ liệu giữa các khối", ["Nguồn", "Đích", "Dữ liệu hoặc lệnh", "Giao thức"], [
        ["DHT11", "ESP32", "Nhiệt độ, độ ẩm không khí", "Tín hiệu số"],
        ["Soil Moisture", "ESP32", "Giá trị ADC độ ẩm đất", "Analog"],
        ["ESP32", "Web", "JSON cảm biến và trạng thái", "HTTP trong LAN"],
        ["Web", "ESP32", "Lệnh relay và cấu hình tự động", "HTTP GET"],
        ["OpenWeather", "Web", "Thời tiết hiện tại và dự báo", "HTTPS API"],
        ["Web", "Gemini", "Ngữ cảnh cảm biến và câu hỏi", "HTTPS qua middleware"],
        ["ESP32", "Relay và bơm", "Bật hoặc tắt tưới", "GPIO active LOW"],
    ], [3.2, 2.6, 5.5, 3.0], small=True)

    add_heading(doc, "3 Thiết kế phần cứng", 1)
    add_heading(doc, "3.1 Sơ đồ kết nối", 2)
    add_table(doc, "Bảng", "Bảng kết nối chân của hệ thống", ["Thiết bị", "Chân thiết bị", "Chân ESP32 hoặc nguồn", "Ghi chú"], [
        ["DHT11", "DATA", "GPIO4", "Dùng thư viện DHT; kiểm tra đọc lỗi."],
        ["DHT11", "VCC, GND", "Nguồn phù hợp, GND", "Theo module thực tế."],
        ["Soil Moisture", "AO", "GPIO34", "ADC1, đầu vào; không có pull-up nội."],
        ["Relay kênh 1", "IN1", "GPIO26", "Active LOW theo firmware."],
        ["Relay", "VCC, GND", "Nguồn relay, GND", "Kiểm tra mức kích của module."],
        ["Bơm MB370", "Hai dây nguồn", "Qua COM và NO", "Dùng nguồn riêng phù hợp với bơm."],
    ], [3.2, 3.0, 4.2, 4.3], small=True)
    add_placeholder(doc, "Hình", "Sơ đồ đấu nối ESP32 cảm biến relay và máy bơm", 8)

    add_heading(doc, "3.2 Thiết kế nguồn và an toàn điện", 2)
    add_body(doc, "Bộ chuyển đổi nguồn phải đáp ứng dòng làm việc và dòng khởi động của bơm. Nguồn logic cần ổn định để ESP32 không khởi động lại khi relay đóng. Nếu dùng một nguồn chung, nhánh bơm nên đi dây riêng từ đầu nguồn và có tụ lọc gần tải; nếu dùng hai nguồn, cần kiểm tra yêu cầu nối chung GND của module relay.")
    add_body(doc, "Trước khi nối bơm, nhóm cần kiểm tra relay bằng tải thử hoặc đồng hồ. Khi kết nối hoàn chỉnh, phải đặt bơm và mạch ở vị trí tránh nước, cố định đầu ống, sử dụng hộp bảo vệ và ngắt nguồn trước khi thay đổi dây. Không chạm vào phần dẫn điện khi hệ thống đang hoạt động.")

    add_heading(doc, "4 Thiết kế thuật toán điều khiển", 1)
    add_heading(doc, "4.1 Chu trình đọc cảm biến", 2)
    add_body(doc, "Firmware gọi server.handleClient liên tục để phục vụ yêu cầu mạng. Sau mỗi 2 giây, ESP32 đọc ADC cảm biến đất, ánh xạ giá trị sang phần trăm, đọc DHT11 và cập nhật bộ nhớ đệm. Nếu DHT11 trả về giá trị không hợp lệ, nhiệt độ và độ ẩm cũ không được coi là dữ liệu mới; API báo trạng thái lỗi để giao diện xử lý.")
    add_code(doc, [
        "#define DHTPIN 4",
        "#define SOIL_MOISTURE_PIN 34",
        "#define RELAY_PIN 26",
        "const int SOIL_DRY_VALUE = 3000;",
        "const int SOIL_WET_VALUE = 1200;",
        "const unsigned long READ_INTERVAL_MS = 2000;",
    ])

    add_heading(doc, "4.2 Quy đổi độ ẩm đất", 2)
    add_body(doc, "Đối với đầu dò điện trở của mô hình, giá trị ADC lớn biểu thị đất khô và giá trị nhỏ biểu thị đất ẩm. Hàm map đảo chiều hai mốc để tạo thang 0 đến 100 phần trăm, sau đó constrain giới hạn kết quả. Việc hiệu chuẩn phải đo mẫu đất khô đại diện và đất đã đạt độ ẩm mục tiêu; không nên dùng mốc mặc định cho mọi loại đất.")

    add_heading(doc, "4.3 Chế độ tưới thủ công", 2)
    add_body(doc, "Khi chế độ tự động tắt, người dùng có thể bật bơm từ giao diện và đặt thời gian chạy. ESP32 lưu thời điểm bật cùng thời lượng yêu cầu, rồi tự tắt khi hết thời gian. Nếu người dùng tắt sớm, relay chuyển về trạng thái OFF ngay. Lệnh thủ công bị từ chối khi chế độ tự động đang bật để tránh hai nguồn điều khiển xung đột.")

    add_heading(doc, "4.4 Chế độ tưới tự động", 2)
    add_body(doc, "Cấu hình tự động gồm enabled, startPercent, stopPercent và durationSec. Khi độ ẩm đất nhỏ hơn hoặc bằng startPercent, ESP32 bật bơm và đánh dấu relay do chế độ tự động kích hoạt. Bơm dừng khi độ ẩm đạt stopPercent hoặc khi hết thời gian tối đa. Hai ngưỡng tạo vùng trễ, hạn chế relay đóng ngắt liên tục quanh một giá trị.")
    add_code(doc, [
        "if (!relayState && soilPercent <= startPercent) setRelay(true, true);",
        "if (relayAutoTriggered && soilPercent >= stopPercent) setRelay(false);",
        "if (relayAutoTriggered && elapsed >= maxDurationMs) setRelay(false);",
    ])
    add_placeholder(doc, "Hình", "Lưu đồ thuật toán đọc cảm biến và tưới tự động", 8)

    add_heading(doc, "4.5 Cơ chế an toàn", 2)
    add_body(doc, "Ngoài thời gian do người dùng cấu hình, firmware có giới hạn cứng RELAY_MAX_ON_MS bằng 10 phút. Giới hạn này áp dụng cho cả điều khiển tay và tự động, giúp giảm rủi ro khi trình duyệt mất kết nối hoặc cấu hình sai. Khi khởi động, firmware gọi setRelay(false) để đảm bảo bơm ở trạng thái tắt trước khi kết nối Wi-Fi.")

    add_heading(doc, "5 Thiết kế giao tiếp ESP32", 1)
    add_table(doc, "Bảng", "Các endpoint HTTP của ESP32", ["API", "Phương\nthức", "Mục đích", "Kết quả chính"], [
        ["/status", "GET", "Xác nhận đúng thiết bị", "device, status, ip, uptime_ms"],
        ["/sensors", "GET", "Lấy cảm biến và trạng thái", "DHT11, soil, relay, auto config"],
        ["/relay", "GET", "Điều khiển bơm thủ công", "relay, manual_duration_sec"],
        ["/auto", "GET", "Đọc hoặc cập nhật tự động", "enabled, start, stop, duration"],
        ["Các endpoint", "OPTIONS", "Hỗ trợ CORS trong LAN", "HTTP 204"],
    ], [3.0, 3.0, 4.5, 4.5], small=True)
    add_body(doc, "Endpoint /status được dùng như bước kiểm tra thiết bị: giao diện chỉ chấp nhận phản hồi có device bằng smart-garden-esp32. Sau khi kết nối, web lấy /sensors ngay và tiếp tục thăm dò định kỳ 5 giây. Mỗi yêu cầu có timeout 3 giây để giao diện không chờ vô hạn.")

    add_heading(doc, "6 Thiết kế ứng dụng web", 1)
    add_heading(doc, "6.1 Giao diện hai khu vực", 2)
    add_body(doc, "Giao diện gồm hai ô chính trên nền gradient xanh. Ô bên trái chứa thanh tìm kiếm khu vực, thời tiết hiện tại và dự báo 5 ngày từ OpenWeather. Khi chưa tìm kiếm, ô chỉ hiển thị thanh nhập và hướng dẫn. Ô bên phải chứa địa chỉ IP ESP32, tên cây, số đo cảm biến, điều khiển bơm, cấu hình tự động và cảnh báo. Trên điện thoại, hai ô xếp dọc và không tràn ngang.")
    add_placeholder(doc, "Hình", "Giao diện tổng thể của ứng dụng Smart Garden", 7)
    add_placeholder(doc, "Hình", "Ô thời tiết OpenWeather sau khi tìm kiếm khu vực", 6)
    add_placeholder(doc, "Hình", "Ô ESP32 hiển thị cảm biến và điều khiển bơm", 6)

    add_heading(doc, "6.2 Quản lý kết nối và trạng thái", 2)
    add_body(doc, "Người dùng chủ động nhập IP và kết nối; ứng dụng không tự dùng lại một thiết bị cũ khi trang vừa mở. Mỗi lần kết nối có session riêng để phản hồi trễ của thiết bị trước không ghi đè dữ liệu mới. Trong lúc gửi lệnh relay hoặc cấu hình, polling tạm dừng để phản hồi cũ không làm mất trạng thái vừa lưu.")

    add_heading(doc, "6.3 Tích hợp thời tiết", 2)
    add_body(doc, "Weather service gửi yêu cầu theo tên khu vực, dùng đơn vị metric và ngôn ngữ tiếng Việt. Dữ liệu hiện tại gồm nhiệt độ, độ ẩm, gió, mưa và mô tả thời tiết; dự báo được tổng hợp theo ngày. Phần phân tích đối chiếu số đo DHT11 tại vườn với OpenWeather khu vực, đồng thời cảnh báo các trường hợp đất khô, mưa hoặc độ ẩm cao.")

    add_heading(doc, "6.4 Tích hợp chatbot gợi ý tưới", 2)
    add_body(doc, "Chatbot nhận tên cây nếu người dùng nhập, số đo độ ẩm đất, dữ liệu DHT11 hợp lệ, trạng thái bơm, cấu hình tự động hiện tại và thời tiết đã tìm kiếm. Phản hồi bắt buộc ở dạng JSON gồm lời giải thích và một recommendation có cấu trúc. Ứng dụng kiểm tra kiểu dữ liệu, giới hạn ngưỡng và thời gian trước khi hiển thị nút thao tác.")
    add_body(doc, "Ba lựa chọn được thiết kế rõ ràng. Áp dụng ngay gửi thông số tới ESP32 và chỉ báo thành công khi thiết bị phản hồi đúng giá trị. Tôi sẽ xem xét điền gợi ý vào form để người dùng sửa trước khi lưu. Từ chối chỉ đóng quyết định của gợi ý và không gửi lệnh. Nếu đổi IP, kết nối lại hoặc đổi tên cây, các nút của gợi ý cũ bị khóa.")
    add_placeholder(doc, "Hình", "Chatbot hiển thị gợi ý tưới và ba lựa chọn", 6)

    add_chapter(doc, "CHƯƠNG 3 TRIỂN KHAI VÀ KIỂM THỬ HỆ THỐNG")
    add_heading(doc, "1 Công cụ và môi trường phát triển", 1)
    add_table(doc, "Bảng", "Công cụ triển khai dự án", ["Thành phần", "Công cụ hoặc công nghệ", "Vai trò"], [
        ["Firmware", "Arduino IDE, C/C++", "Biên dịch và nạp chương trình cho ESP32."],
        ["Thư viện ESP32", "WiFi, WebServer, ESPmDNS", "Kết nối mạng, HTTP và mDNS."],
        ["Cảm biến", "Adafruit Sensor, DHT", "Đọc nhiệt độ và độ ẩm từ DHT11."],
        ["Giao diện", "React 18, Vite, Tailwind CSS", "Xây dựng và đóng gói ứng dụng web."],
        ["Biểu tượng", "lucide-react", "Hiển thị biểu tượng giao diện."],
        ["Dịch vụ", "OpenWeather, Gemini", "Dữ liệu thời tiết và gợi ý tưới."],
    ], [3.5, 5.0, 6.5], small=True)

    add_heading(doc, "2 Triển khai firmware ESP32", 1)
    add_heading(doc, "2.1 Khởi tạo thiết bị", 2)
    add_body(doc, "Trong setup, firmware khởi tạo Serial ở 115200 baud, DHT11, chân relay và trạng thái an toàn. ESP32 kết nối Wi-Fi bằng thông tin trong arduino_secret.h, sau đó khởi động mDNS với tên smartgarden và đăng ký các endpoint. Giá trị cảm biến đầu tiên được đọc sau khi web server sẵn sàng.")
    add_code(doc, [
        "WiFi.mode(WIFI_STA);",
        "WiFi.begin(WIFI_SSID, WIFI_PASSWORD);",
        "MDNS.begin(\"smartgarden\");",
        "server.on(\"/status\", HTTP_GET, handleStatus);",
        "server.on(\"/sensors\", HTTP_GET, handleSensors);",
        "server.on(\"/relay\", HTTP_GET, handleRelay);",
        "server.on(\"/auto\", HTTP_GET, handleAuto);",
    ])

    add_heading(doc, "2.2 Đọc và đóng gói dữ liệu", 2)
    add_body(doc, "SensorCache lưu nhiệt độ, độ ẩm, giá trị ADC, phần trăm độ ẩm đất và thời điểm cập nhật. Handler /sensors đóng gói dữ liệu cùng relay và cấu hình tự động thành JSON. Cách tổ chức này tách thao tác đọc phần cứng khỏi yêu cầu mạng, giúp mỗi phản hồi nhanh và nhất quán.")

    add_heading(doc, "2.3 Điều khiển relay", 2)
    add_body(doc, "Hàm setRelay là điểm điều khiển duy nhất của relay. Khi bật, hàm ghi nhận thời điểm và nguồn kích hoạt; khi tắt, cờ tự động được xóa. Do relay active LOW, hàm xuất LOW khi bật và HIGH khi tắt. Việc gom logic vào một hàm giảm nguy cơ trạng thái phần mềm lệch với chân điều khiển.")

    add_heading(doc, "3 Triển khai ứng dụng web", 1)
    add_heading(doc, "3.1 Quản lý thời tiết", 2)
    add_body(doc, "Hook useWeather quản lý dữ liệu hiện tại, dự báo, cảnh báo, tải và lỗi. WeatherPanel chỉ hiển thị lời nhắc trước lần tìm kiếm; sau khi có dữ liệu, giao diện trình bày vị trí, nhiệt độ, độ ẩm, gió, mưa, bình minh, hoàng hôn và năm ngày tiếp theo.")

    add_heading(doc, "3.2 Quản lý cảm biến", 2)
    add_body(doc, "Hook useSensor quản lý địa chỉ nhập, thiết bị đang hoạt động, phiên kết nối, số đo, trạng thái relay và thao tác cấu hình. Polling dùng setTimeout nối tiếp thay vì setInterval để tránh chồng yêu cầu. Khi lệnh thay đổi thiết bị đang chạy, polling dừng và chỉ khởi động lại sau khi lệnh hoàn thành.")

    add_heading(doc, "3.3 Kiểm tra cấu hình tưới", 2)
    add_body(doc, "Trước khi gửi, ứng dụng yêu cầu enabled là Boolean; startPercent, stopPercent và durationSec là số nguyên; start nhỏ hơn stop; thời gian từ 1 đến 600 giây. Sau khi ESP32 phản hồi, ứng dụng so sánh toàn bộ trường với yêu cầu. Nếu giá trị không khớp, giao diện báo chưa xác nhận đúng cấu hình thay vì hiển thị thành công.")

    add_heading(doc, "3.4 Kết nối Gemini", 2)
    add_body(doc, "Yêu cầu Gemini đi qua middleware của Vite tới endpoint v1beta. Middleware chỉ nhận đúng đường dẫn tạo nội dung, chuyển tiếp nội dung JSON và luôn trả lỗi có cấu trúc khi upstream rỗng, lỗi mạng hoặc timeout. Ứng dụng tiếp tục kiểm tra JSON của mô hình trước khi tạo các nút áp dụng cấu hình.")

    add_heading(doc, "4 Quy trình lắp ráp và vận hành", 1)
    add_bullets(doc, [
        "Kiểm tra điện áp đầu ra của bộ chuyển đổi trước khi gắn ESP32, relay và bơm.",
        "Đấu DHT11 vào GPIO4, Soil Moisture vào GPIO34 và relay kênh 1 vào GPIO26 theo sơ đồ.",
        "Đấu bơm qua tiếp điểm COM và NO, cố định ống hút và ống ra, đặt mạch xa nguồn nước.",
        "Nạp firmware, mở Serial Monitor và ghi lại địa chỉ IP do ESP32 nhận từ router.",
        "Chạy ứng dụng web, nhập IP, kết nối và xác nhận ba số đo cảm biến hợp lý.",
        "Hiệu chuẩn mốc khô và ướt của cảm biến đất bằng mẫu đất thực tế.",
        "Thử bơm thủ công với thời gian ngắn, quan sát lưu lượng và kiểm tra rò rỉ.",
        "Thiết lập ngưỡng tự động, thử các trạng thái đất khô và đất đủ ẩm, xác nhận bơm dừng đúng điều kiện.",
    ])
    add_placeholder(doc, "Hình", "Mô hình phần cứng sau khi lắp ráp", 8)
    add_placeholder(doc, "Hình", "Vị trí cảm biến đất DHT11 và đầu tưới trên mô hình", 7)

    add_heading(doc, "5 Kế hoạch kiểm thử", 1)
    add_table(doc, "Bảng", "Các ca kiểm thử chức năng", ["Mã", "Điều kiện", "Thao tác", "Kết quả mong đợi"], [
        ["01", "ESP32 và web cùng Wi-Fi", "Nhập IP và kết nối", "Nhận diện đúng smart-garden-esp32 và hiển thị số đo."],
        ["02", "Đất khô", "Đọc /sensors", "soil_percent giảm về vùng khô sau hiệu chuẩn."],
        ["03", "DHT11 hoạt động", "Theo dõi nhiều chu kỳ", "Nhiệt độ, độ ẩm cập nhật; dht11_ok bằng true."],
        ["04", "Auto tắt", "Bật bơm 30 giây", "Relay bật, hiển thị thủ công và tự tắt đúng thời gian."],
        ["05", "Auto bật, đất dưới ngưỡng", "Chờ chu kỳ đọc", "Bơm bật và relay_auto_triggered bằng true."],
        ["06", "Đang tưới tự động", "Làm đất đạt ngưỡng dừng", "Bơm tắt khi soil_percent đạt stopPercent."],
        ["07", "DHT11 ngắt", "Đọc /sensors", "dht11_ok bằng false; web không hiển thị số giả."],
        ["08", "Mất Wi-Fi", "Chờ polling", "Web báo mất kết nối; nút điều khiển bị khóa."],
        ["09", "Chatbot có cảm biến", "Yêu cầu gợi ý tưới", "Hiển thị lời giải thích, thông số hợp lệ và ba nút."],
        ["10", "Gợi ý cũ", "Đổi tên cây hoặc kết nối", "Áp dụng và xem xét bị khóa; yêu cầu gợi ý mới."],
    ], [1.6, 3.2, 3.7, 6.0], small=True)

    add_heading(doc, "6 Kết quả kiểm thử phần mềm", 1)
    add_body(doc, "Tại thời điểm hoàn thiện tài liệu, bộ kiểm thử tự động của ứng dụng web gồm 17 phép kiểm tra và đã chạy đạt. Các kiểm tra bao phủ dữ liệu AI có cấu trúc, giới hạn cấu hình, phản hồi rỗng hoặc sai định dạng, DHT11 lỗi, payload tới ESP32, xác nhận cấu hình, middleware Gemini và các lỗi mạng mô phỏng. Lệnh build sản phẩm cũng hoàn tất thành công.")
    add_table(doc, "Bảng", "Tóm tắt kết quả kiểm thử phần mềm", ["Nhóm kiểm tra", "Kết quả", "Ghi chú"], [
        ["Kiểm tra cấu hình tưới", "Đạt", "Chặn rỗng, số thực, đảo ngưỡng và thời gian ngoài giới hạn."],
        ["Kiểm tra dữ liệu chatbot", "Đạt", "Chỉ tạo thao tác khi JSON và recommendation hợp lệ."],
        ["Kiểm tra sensor service", "Đạt", "Gửi đúng host, endpoint và yêu cầu ESP32 xác nhận."],
        ["Kiểm tra Gemini middleware", "Đạt", "Đúng đường dẫn, xử lý API lỗi, rỗng, timeout và mất mạng."],
        ["Build giao diện", "Đạt", "Vite tạo gói sản phẩm thành công."],
    ], [5.0, 2.3, 7.2], small=True)

    add_heading(doc, "7 Biểu mẫu ghi nhận kiểm thử phần cứng", 1)
    add_body(doc, "Bảng dưới đây được để trống phần số liệu để nhóm cập nhật sau khi lắp ráp và đo trên mô hình thật. Không nên dùng số liệu giả thay cho kết quả đo. Mỗi ca thử cần ghi ngày, điều kiện, giá trị ADC, phần trăm độ ẩm, nhiệt độ, độ ẩm không khí, thời gian bơm và nhận xét.")
    add_table(doc, "Bảng", "Phiếu ghi nhận số liệu thử nghiệm thực tế", ["STT", "Đất", "ADC", "Ẩm đất\n(%)", "Nhiệt\n(°C)", "Ẩm KK\n(%)", "Bơm\n(s)", "Nhận xét"], [
        [str(i), "", "", "", "", "", "", ""] for i in range(1, 7)
    ], [1.2, 2.3, 1.5, 2.1, 2.0, 2.0, 1.5, 2.6], small=True)
    add_placeholder(doc, "Hình", "Thử nghiệm cảm biến độ ẩm đất ở trạng thái khô", 5)
    add_placeholder(doc, "Hình", "Thử nghiệm cảm biến độ ẩm đất sau khi tưới", 5)
    add_placeholder(doc, "Hình", "Thử nghiệm máy bơm MB370 và relay", 5)

    add_chapter(doc, "CHƯƠNG 4 KẾT QUẢ ĐÁNH GIÁ VÀ HƯỚNG PHÁT TRIỂN")
    add_heading(doc, "1 Những nội dung đã thực hiện", 1)
    add_body(doc, "Dự án đã hình thành đầy đủ kiến trúc phần mềm cho mô hình vườn thông minh: firmware đọc cảm biến và điều khiển bơm; API HTTP trao đổi trạng thái; ứng dụng web hiển thị thời tiết và dữ liệu tại vườn; chatbot tạo gợi ý có cấu trúc; thao tác áp dụng được xác nhận lại từ ESP32. Giao diện được tổ chức thành hai khu vực rõ ràng và hỗ trợ màn hình di động.")
    add_bullets(doc, [
        "Đọc độ ẩm đất theo chu kỳ và quy đổi bằng hai mốc hiệu chuẩn.",
        "Đọc DHT11 và phân biệt dữ liệu hợp lệ với lỗi cảm biến.",
        "Điều khiển bơm thủ công có thời gian và tưới tự động theo hai ngưỡng.",
        "Giới hạn relay 10 phút, tắt bơm khi khởi động và tránh xung đột tay/tự động.",
        "Kết nối ESP32 bằng IP hoặc mDNS và cung cấp dữ liệu JSON trong LAN.",
        "Tích hợp OpenWeather, cảnh báo môi trường và đối chiếu số đo tại vườn.",
        "Tích hợp Gemini với ba lựa chọn áp dụng, xem xét hoặc từ chối.",
        "Kiểm tra dữ liệu đầu vào, phản hồi trễ, thay đổi phiên kết nối và lỗi dịch vụ.",
    ])

    add_heading(doc, "2 Đánh giá theo kịch bản vận hành", 1)
    add_table(doc, "Bảng", "Hành vi hệ thống theo trạng thái", ["Trạng thái", "Quyết định", "Hiển thị và bảo vệ"], [
        ["Đất dưới ngưỡng bắt đầu", "Bật bơm nếu tự động đang bật", "Đánh dấu tưới tự động và theo dõi thời gian."],
        ["Đất đạt ngưỡng dừng", "Tắt bơm", "Cập nhật relay OFF trên web."],
        ["Chưa đạt ngưỡng nhưng hết thời gian", "Tắt bơm", "Ngăn tưới kéo dài."],
        ["Lệnh tay khi auto bật", "Từ chối", "Yêu cầu tắt auto trước khi điều khiển tay."],
        ["DHT11 lỗi", "Không dùng nhiệt độ và độ ẩm", "Hiển thị dấu gạch và cảnh báo kiểm tra cảm biến."],
        ["Mất kết nối ESP32", "Dừng thao tác từ web", "Khóa nút và báo lỗi mạng."],
        ["Gợi ý AI không hợp lệ", "Không tạo nút áp dụng", "Yêu cầu người dùng gợi ý lại."],
    ], [4.1, 4.4, 6.0], small=True)

    add_heading(doc, "3 Hạn chế hiện tại", 1)
    add_body(doc, "Cảm biến đất điện trở cần hiệu chuẩn và có nguy cơ ăn mòn. DHT11 chỉ phù hợp giám sát cơ bản. Hệ thống chưa đo lưu lượng nên thời gian tưới chưa quy đổi trực tiếp thành thể tích nước. Firmware chưa lưu cấu hình tự động vào bộ nhớ không mất dữ liệu; sau khi khởi động lại, cấu hình trở về mặc định.")
    add_body(doc, "Kết nối ESP32 hiện dùng HTTP trong mạng LAN. Nếu giao diện được triển khai trên HTTPS công khai, trình duyệt có thể chặn truy cập tới địa chỉ HTTP nội bộ do mixed content. Khóa OpenWeather và Gemini cần được quản lý phía máy chủ khi triển khai thực tế; không nên đưa khóa bí mật vào gói JavaScript công khai.")
    add_body(doc, "Chatbot chỉ cung cấp khuyến nghị tham khảo. Loại đất, lưu lượng bơm, kích thước chậu, giai đoạn sinh trưởng và điều kiện ánh sáng chưa được đo đầy đủ, vì vậy người dùng vẫn phải theo dõi cây và hiệu chỉnh ngưỡng sau mỗi lần thử.")

    add_heading(doc, "4 Khó khăn và hướng xử lý", 1)
    add_table(doc, "Bảng", "Khó khăn kỹ thuật và giải pháp", ["Khó khăn", "Nguyên nhân", "Giải pháp đã áp dụng hoặc đề xuất"], [
        ["ADC dao động", "Đất không đồng nhất, nguồn và nhiễu", "Hiệu chuẩn tại chỗ, đọc nhiều mẫu và lọc trung bình trong phiên bản sau."],
        ["ESP32 khởi động lại khi bơm chạy", "Sụt áp và dòng khởi động", "Tách nhánh nguồn, chọn bộ chuyển đổi đủ dòng, bổ sung tụ lọc."],
        ["DHT11 đọc lỗi", "Chu kỳ đọc quá nhanh hoặc dây nối", "Đọc mỗi 2 giây, kiểm tra NaN và báo lỗi trên web."],
        ["Phản hồi polling ghi đè lệnh", "Yêu cầu mạng bất đồng bộ", "Dừng polling khi thay đổi trạng thái và dùng phiên kết nối."],
        ["Gemini trả về dữ liệu không chuẩn", "Phản hồi rỗng hoặc sai JSON", "Middleware và client kiểm tra định dạng, giới hạn và lỗi."],
        ["Giao diện HTTPS không gọi được ESP32 HTTP", "Chính sách mixed content", "Dùng gateway cục bộ, cloud API hoặc TLS phù hợp khi triển khai."],
    ], [4.1, 4.5, 6.0], small=True)

    add_heading(doc, "5 Hướng phát triển", 1)
    add_bullets(doc, [
        "Thay đầu dò điện trở bằng cảm biến độ ẩm đất điện dung bền hơn.",
        "Lưu cấu hình bằng Preferences hoặc NVS để giữ ngưỡng sau khi mất điện.",
        "Bổ sung cảm biến mực nước và lưu lượng để bảo vệ bơm chạy khô và đo thể tích tưới.",
        "Điều khiển nhiều vùng bằng kênh relay còn lại hoặc module mở rộng, mỗi vùng có cảm biến riêng.",
        "Ghi lịch sử số đo, biểu đồ xu hướng và nhật ký mỗi lần bơm hoạt động.",
        "Dùng MQTT hoặc gateway đám mây để giám sát ngoài mạng nội bộ với xác thực và mã hóa.",
        "Kết hợp dự báo mưa, loại cây, giai đoạn sinh trưởng và dữ liệu lịch sử để tối ưu lịch tưới.",
        "Thiết kế PCB, hộp chống ẩm và đầu nối tiêu chuẩn để nâng độ bền của mô hình.",
    ])

    add_chapter(doc, "KẾT LUẬN")
    add_body(doc, "Đề tài đã xây dựng được một thiết kế hoàn chỉnh cho mô hình vườn thông minh sử dụng ESP32, DHT11, cảm biến độ ẩm đất, máy bơm nước MB370, relay 2 kênh, mạch mở rộng và bộ chuyển đổi nguồn. Firmware tổ chức rõ luồng đọc cảm biến, điều khiển bơm và cung cấp API; ứng dụng web thống nhất dữ liệu thời tiết, trạng thái tại vườn và cấu hình tưới trong một giao diện hai khu vực.")
    add_body(doc, "Điểm quan trọng của hệ thống là người dùng vẫn kiểm soát quyết định tưới. Chế độ tự động có vùng trễ và giới hạn thời gian; chatbot không tự thay đổi thiết bị mà yêu cầu người dùng áp dụng, xem xét hoặc từ chối. Các kiểm tra phần mềm đã xác nhận đường dẫn dữ liệu và xử lý lỗi. Sau khi bổ sung ảnh cùng số liệu thử nghiệm thực tế, báo cáo có thể dùng để trình bày đầy đủ quá trình lắp ráp và đánh giá mô hình.")
    add_body(doc, "Trong giai đoạn tiếp theo, nhóm cần tập trung hiệu chuẩn cảm biến đất, đo lưu lượng bơm, hoàn thiện bảo vệ nguồn và ghi nhận kết quả theo các ca kiểm thử đã xây dựng. Đây là cơ sở để chuyển mô hình học tập thành một hệ thống tưới ổn định hơn, có khả năng lưu trữ lịch sử và quản lý nhiều vùng cây.")

    add_chapter(doc, "TÀI LIỆU THAM KHẢO")
    references = [
        "[1] Espressif Systems, ESP32 Datasheet và ESP-IDF Programming Guide, https://docs.espressif.com/.",
        "[2] Espressif Systems, GPIO & RTC GPIO - ESP32, https://docs.espressif.com/projects/esp-idf/en/latest/esp32/api-reference/peripherals/gpio.html.",
        "[3] Espressif Systems, Analog to Digital Converter - ESP32, https://docs.espressif.com/projects/esp-idf/en/latest/esp32/api-reference/peripherals/adc_oneshot.html.",
        "[4] Aosong Electronics, DHT11 Temperature and Humidity Sensor Datasheet, https://www.aosong.com/.",
        "[5] OpenWeather, Current Weather Data và 5 Day Weather Forecast, https://openweathermap.org/api.",
        "[6] Google AI for Developers, Gemini API Generate Content, https://ai.google.dev/api/generate-content.",
        "[7] React Documentation, https://react.dev/.",
        "[8] Vite Documentation, https://vite.dev/.",
        "[9] Nhóm 6, mã nguồn firmware smart_garden_phase3.ino và ứng dụng Smart Garden trong thư mục dự án, 2026.",
    ]
    for ref in references:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(-0.7)
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(5)
        set_font(p.add_run(ref), size=12)

    add_chapter(doc, "PHỤ LỤC A THÔNG SỐ CẤU HÌNH HIỆN TẠI")
    add_table(doc, "Bảng", "Thông số mặc định trong firmware", ["Tham số", "Giá trị", "Ý nghĩa"], [
        ["SOIL_DRY_VALUE", "3000", "Mốc ADC đất khô, cần hiệu chuẩn lại."],
        ["SOIL_WET_VALUE", "1200", "Mốc ADC đất ướt, cần hiệu chuẩn lại."],
        ["startPercent", "30%", "Bật tưới tự động khi độ ẩm nhỏ hơn hoặc bằng."],
        ["stopPercent", "70%", "Tắt tưới tự động khi độ ẩm lớn hơn hoặc bằng."],
        ["maxDuration", "120 giây", "Thời gian tối đa mặc định của một lần tưới auto."],
        ["RELAY_MAX_ON", "10 phút", "Giới hạn cứng cho mọi chế độ."],
        ["READ_INTERVAL", "2 giây", "Chu kỳ đọc cảm biến trên ESP32."],
        ["WEB_POLL_INTERVAL", "5 giây", "Chu kỳ lấy dữ liệu tại giao diện."],
        ["REQUEST_TIMEOUT", "3 giây", "Thời gian chờ yêu cầu ESP32."],
    ], [5.0, 3.1, 6.4], small=True)

    add_chapter(doc, "PHỤ LỤC B DANH SÁCH ẢNH CẦN BỔ SUNG")
    add_body(doc, "Khi có ảnh thực tế, thay trực tiếp từng khung ảnh tương ứng và giữ lại dòng chú thích. Ảnh nên đủ sáng, nền gọn, không để lộ mật khẩu Wi-Fi hoặc khóa API. Với ảnh giao diện, che địa chỉ IP nếu tài liệu được công khai.")
    add_table(doc, "Bảng", "Danh sách ảnh cần chuẩn bị", ["STT", "Ảnh cần chụp", "Yêu cầu"], [
        ["1", "ESP32", "Ảnh rõ tên board và các chân đang dùng."],
        ["2", "DHT11", "Ảnh cảm biến và vị trí lắp trên mô hình."],
        ["3", "Soil Moisture", "Ảnh đầu dò trong đất và module xử lý."],
        ["4", "Máy bơm MB370", "Ảnh bơm, nguồn và đường ống."],
        ["5", "Relay 2 kênh", "Ảnh rõ COM, NO và dây điều khiển."],
        ["6", "Mạch mở rộng và bộ chuyển đổi", "Ảnh toàn bộ dây nguồn và đầu ra điện áp."],
        ["7", "Mô hình hoàn chỉnh", "Ảnh tổng thể nhìn được cây, bình nước và mạch."],
        ["8", "Giao diện web", "Ảnh ô thời tiết, ô ESP32 và chatbot."],
        ["9", "Kiểm thử", "Ảnh đất khô, đất ẩm, bơm chạy và bơm dừng."],
    ], [1.2, 5.2, 8.1], small=True)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    core = doc.core_properties
    core.title = "Xây dựng mô hình vườn thông minh có cảm biến đo độ ẩm đất nhiệt độ độ ẩm có sử dụng ESP32"
    core.subject = "Báo cáo bài tập lớn môn Phát triển ứng dụng hệ thống nhúng và IoT"
    core.author = "Nhóm 6 - DH13C6"
    core.keywords = "ESP32, Smart Garden, DHT11, Soil Moisture, MB370, IoT"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
