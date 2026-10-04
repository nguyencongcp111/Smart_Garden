from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
import json

src = Path(__file__).with_name("reference.docx")
doc = Document(src)

def text_of_paragraph(p):
    return " ".join(p.text.split())

headings = []
nonempty = []
for i, p in enumerate(doc.paragraphs):
    text = text_of_paragraph(p)
    if not text:
        continue
    nonempty.append({"index": i, "style": p.style.name if p.style else "", "text": text})
    if p.style and p.style.name.startswith("Heading"):
        headings.append(nonempty[-1])

tables = []
for ti, table in enumerate(doc.tables):
    rows = []
    for row in table.rows:
        rows.append([" ".join(cell.text.split()) for cell in row.cells])
    tables.append({"index": ti, "rows": rows})

styles = {}
for p in doc.paragraphs:
    if not p.style:
        continue
    name = p.style.name
    if name in styles:
        continue
    style = p.style
    font = style.font
    pf = style.paragraph_format
    styles[name] = {
        "font": font.name,
        "size_pt": font.size.pt if font.size else None,
        "bold": font.bold,
        "italic": font.italic,
        "alignment": str(pf.alignment),
        "space_before_pt": pf.space_before.pt if pf.space_before else None,
        "space_after_pt": pf.space_after.pt if pf.space_after else None,
        "line_spacing": str(pf.line_spacing),
    }

headers = []
footers = []
for si, section in enumerate(doc.sections):
    headers.append({"section": si + 1, "text": " | ".join(text_of_paragraph(p) for p in section.header.paragraphs if text_of_paragraph(p))})
    footers.append({"section": si + 1, "text": " | ".join(text_of_paragraph(p) for p in section.footer.paragraphs if text_of_paragraph(p))})

payload = {
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "headings": headings,
    "nonempty_paragraphs": nonempty,
    "tables": tables,
    "styles": styles,
    "headers": headers,
    "footers": footers,
}
print(json.dumps(payload, ensure_ascii=False, indent=2))
